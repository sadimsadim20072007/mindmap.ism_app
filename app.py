import streamlit as st
import google.generativeai as genai
import pypdf
from pptx import Presentation

# إعدادات الصفحة
st.set_page_config(page_title="صانع الخرائط الذهنية المرئية", page_icon="🧠", layout="wide")

st.title("🧠 صانع الخرائط الذهنية بالهيكلة المحددة")
st.write("ارفع ملفك (PDF, TXT, PPTX) وسيقوم التطبيق بتوليد خطاطة مرئية تلتزم بنمط البطاقة التربوية تماماً.")

# مفتاح API
api_key = st.sidebar.text_input("أدخل مفتاح Gemini API الخاص بك:", type="password")

# رفع الملفات
uploaded_file = st.file_uploader("اختر ملفاً لرفعه", type=["pdf", "txt", "pptx"])

text_content = ""

if uploaded_file is not None:
    if uploaded_file.name.endswith(".pdf"):
        pdf_reader = pypdf.PdfReader(uploaded_file)
        for page in pdf_reader.pages:
            text_content += page.extract_text() or ""
    elif uploaded_file.name.endswith(".txt"):
        text_content = uploaded_file.read().decode("utf-8")
    elif uploaded_file.name.endswith(".pptx"):
        prs = Presentation(uploaded_file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text_content += paragraph.text + "\n"

    st.success("تم قراءة الملف بنجاح!")

    if st.button("🚀 أنشئ الخطاطة المرئية"):
        if not api_key:
            st.error("يرجى إدخال مفتاح API في القائمة الجانبية أولاً.")
        else:
            try:
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel("gemini-3.8-flash")
                
                with st.spinner("جاري تصميم الخطاطة بالهيكلة المحددة..."):
                    prompt = f"""
                    قم بتحليل النص التالي واستخراج خطاطة ذهنية تتبع **بالضبط وبدون أي تغيير** الهيكلة التالية للنموذج المرفق:

                    1. **العقدة الرئيسية (الموضوع / العنوان)**
                    2. **الفروع الرئيسية الأربعة الإلزامية:**
                       - **الأهداف والمعلومات العامة**
                       - **المراحل والأنشطة الزمنية** (افتتاح الحصة، أنشطة القراءة، أنشطة الكتابة، الاختتام)
                       - **الوسائل والدعامات التربوية**
                       - **التقويم والواجبات المنزلية**

                    قم بصياغة النتيجة كـ كود Graphviz DOT فقط (دون أي نص خارجي)، ابدأ بـ `digraph G {{` وانته بـ `}}`.
                    استخدم الاتجاه الأفقي `rankdir=LR;` وأشكال مستطيلات ملونة لكل فرع.

                    النص المراد تحليله:
                    \n\n{text_content[:8000]}
                    """
                    
                    response = model.generate_content(prompt)
                    dot_code = response.text.strip()
                    
                    if "```" in dot_code:
                        lines = dot_code.split("\n")
                        dot_code = "\n".join([line for line in lines if not line.startswith("```")])
                    
                    st.subheader("📌 الخطاطة الذهنية بالهيكلة المحددة:")
                    st.graphviz_chart(dot_code)
                    
            except Exception as e:
                st.error(f"حدث خطأ أثناء رسم الخطاطة: {e}")
